from docx import Document
d = Document(r"D:\antigravity\longpost\.verify\test.docx")
print("paragraphs:", len(d.paragraphs))
for p in d.paragraphs[:6]:
    print(" -", repr(p.style.name), "|", p.text[:40])
