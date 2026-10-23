from pptx import Presentation

path = r"C:\Users\14948\Documents\xwechat_files\wxid_yp3lr7m1hzm122_f948\msg\file\2026-10\LongPost_v3_2_3_Whitepaper.pptx"
prs = Presentation(path)
print("slides:", len(prs.slides), "size:", round(prs.slide_width/914400,2), "x", round(prs.slide_height/914400,2))

def color_of(shape):
    try:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    c = r.font.color
                    if c and c.type is not None:
                        return str(c.rgb)
    except Exception:
        pass
    return None

for i, slide in enumerate(prs.slides):
    if i >= 5:
        break
    print("--- slide", i + 1)
    try:
        fill = slide.background.fill
        if fill.type is not None:
            print("  bg:", fill.fore_color.rgb)
    except Exception:
        pass
    for sh in slide.shapes:
        t = sh.text_frame.text.replace(chr(10), " / ")[:60] if sh.has_text_frame else ""
        print("  ", round(sh.left/914400,2), round(sh.top/914400,2), round(sh.width/914400,2), round(sh.height/914400,2), "|", color_of(sh), "|", t)
